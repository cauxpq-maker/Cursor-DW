# -*- coding: utf-8 -*-
"""把 report.html 转为可在 Word 中打开、编辑的 .docx。"""
import os
import re

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "report.html")
OUT = os.path.join(HERE, "旅顺口区海产品养殖加工产业专题分析报告（完整版）.docx")

TEAL = RGBColor(15, 118, 110)
NAVY = RGBColor(4, 47, 46)
GRAY = RGBColor(71, 85, 105)
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="CBD5E1"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_paragraph_spacing(p, after=6, before=0, line=1.35, first_line=None):
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)


def add_page_break(doc):
    doc.add_page_break()


def text_of(el):
    return re.sub(r"\s+", " ", el.get_text(" ", strip=True)).strip()


def add_runs_from_html(paragraph, node, size=11, color=None):
    """把节点内的文本/加粗写进段落。"""
    if node is None:
        return
    if isinstance(node, NavigableString):
        t = str(node)
        if t.strip() == "" and t != "":
            return
        if t:
            run = paragraph.add_run(t)
            set_run_font(run, size=size, color=color)
        return
    if not isinstance(node, Tag):
        return
    if node.name in ("br",):
        paragraph.add_run().add_break()
        return
    if node.name in ("b", "strong"):
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=size, bold=True, color=color or NAVY)
        return
    if node.name in ("span", "i", "em"):
        for child in node.children:
            add_runs_from_html(paragraph, child, size=size, color=color)
        return
    for child in node.children:
        add_runs_from_html(paragraph, child, size=size, color=color)


def shade_block(doc, html_node, fill, left_border_color, size=10.5):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, left_border_color)
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2, line=1.3)
    add_runs_from_html(p, html_node, size=size)
    # extra paragraphs for nested p/li
    first = True
    blocks = html_node.find_all(["p", "li"], recursive=True)
    if html_node.name in ("p",) or not blocks:
        pass
    else:
        p.clear()
        for i, block in enumerate(blocks):
            para = p if i == 0 else cell.add_paragraph()
            set_paragraph_spacing(para, after=3, line=1.3)
            prefix = ""
            if block.name == "li":
                parent = block.parent
                if parent and parent.name == "ol":
                    prefix = f"{len(list(block.find_previous_siblings('li'))) + 1}. "
                else:
                    prefix = "• "
            if prefix:
                run = para.add_run(prefix)
                set_run_font(run, size=size, bold=True, color=TEAL)
            add_runs_from_html(para, block, size=size)
    doc.add_paragraph()


def add_table(doc, table_el):
    rows = table_el.find_all("tr")
    if not rows:
        return
    ncols = max(len(r.find_all(["th", "td"])) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.autofit = True
    for r_i, tr in enumerate(rows):
        cells = tr.find_all(["th", "td"])
        is_header = all(c.name == "th" for c in cells) or tr.find("th")
        for c_i in range(ncols):
            cell = tbl.cell(r_i, c_i)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=2, line=1.2)
            if c_i < len(cells):
                add_runs_from_html(
                    p,
                    cells[c_i],
                    size=9.5,
                    color=WHITE if (r_i == 0 and is_header) else NAVY if "hd" in (cells[c_i].get("class") or []) else None,
                )
                for run in p.runs:
                    if r_i == 0 and is_header:
                        run.bold = True
                        run.font.color.rgb = WHITE
            if r_i == 0 and is_header:
                set_cell_shading(cell, "042F2E")
            elif r_i % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.add_paragraph()


def add_image(doc, src, caption=None):
    path = os.path.join(HERE, src)
    if not os.path.isfile(path):
        p = doc.add_paragraph(f"[缺图：{src}]")
        set_run_font(p.runs[0], size=10, color=GRAY)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=2)
    run = p.add_run()
    run.add_picture(path, width=Cm(15.6))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(cap, after=8)
        r = cap.add_run(caption)
        set_run_font(r, size=9, color=GRAY)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.line_spacing = 1.35

    for style_name, size, color in (
        ("Heading 1", 16, NAVY),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 12, RGBColor(17, 94, 89)),
        ("Title", 22, NAVY),
        ("Subtitle", 13, TEAL),
    ):
        st = doc.styles[style_name]
        st.font.name = "微软雅黑"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_cover(doc, cover):
    tag = cover.find(class_="tag")
    h1 = cover.find("h1")
    h2 = cover.find("h2")
    points = cover.select(".points div")
    meta = cover.select(".meta div")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=24, after=12)
    r = p.add_run(text_of(tag) if tag else "产业研究 · 海洋经济专题")
    set_run_font(r, size=12, color=TEAL)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, after=8)
    t = text_of(h1).replace("产业专题分析报告", "\n产业专题分析报告") if h1 else ""
    r = title.add_run("旅顺口区海产品养殖加工产业专题分析报告")
    set_run_font(r, size=22, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(sub, after=16)
    r = sub.add_run(text_of(h2).replace("专项 ", "专项\n") if h2 else "")
    set_run_font(r, size=12, color=TEAL)

    for div in points:
        line = text_of(div).lstrip("◆ ").strip()
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=4, line=1.4)
        r = p.add_run("◆  " + line)
        set_run_font(r, size=11, color=NAVY)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=18, after=4)
    r = p.add_run("—" * 28)
    set_run_font(r, size=10, color=GRAY)
    for div in meta:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=2)
        r = p.add_run(text_of(div))
        set_run_font(r, size=10, color=GRAY)

    note = doc.add_paragraph()
    set_paragraph_spacing(note, before=14, after=6)
    r = note.add_run("说明：本文件为可编辑 Word 稿。标题可在“导航窗格”中跳转，表格可直接改单元格，图为嵌入图片，可替换或删改。")
    set_run_font(r, size=10, color=GRAY)

    add_page_break(doc)


def convert():
    with open(SRC, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "lxml")
    body = soup.body
    doc = Document()
    configure_styles(doc)

    cover = body.find(class_="cover")
    if cover:
        add_cover(doc, cover)
        cover.decompose()

    for el in list(body.children):
        if isinstance(el, NavigableString):
            continue
        if not isinstance(el, Tag):
            continue
        classes = el.get("class") or []

        if el.name == "h1":
            p = doc.add_heading(text_of(el), level=1)
            for run in p.runs:
                set_run_font(run, size=16, bold=True, color=NAVY)
            continue
        if el.name == "h2":
            p = doc.add_heading(text_of(el), level=2)
            for run in p.runs:
                set_run_font(run, size=13, bold=True, color=TEAL)
            continue
        if el.name == "h3":
            p = doc.add_heading(text_of(el), level=3)
            for run in p.runs:
                set_run_font(run, size=12, bold=True, color=RGBColor(17, 94, 89))
            continue
        if el.name == "p":
            p = doc.add_paragraph()
            indent = 0.74 if "note" not in classes else 0
            set_paragraph_spacing(p, after=6, first_line=indent)
            add_runs_from_html(p, el, size=10 if "note" in classes else 11,
                               color=GRAY if "note" in classes else None)
            continue
        if el.name in ("ul", "ol"):
            ordered = el.name == "ol"
            for i, li in enumerate(el.find_all("li", recursive=False), 1):
                p = doc.add_paragraph()
                set_paragraph_spacing(p, after=3, line=1.35)
                mark = f"{i}. " if ordered else "• "
                r = p.add_run(mark)
                set_run_font(r, size=11, bold=True, color=TEAL)
                add_runs_from_html(p, li, size=11)
            continue
        if el.name == "table":
            add_table(doc, el)
            continue
        if el.name == "figure":
            img = el.find("img")
            cap = el.find("figcaption")
            add_image(doc, img.get("src") if img else "", text_of(cap) if cap else None)
            continue
        if "warnbox" in classes:
            shade_block(doc, el, "FEF2F2", "DC2626")
            continue
        if "okbox" in classes:
            shade_block(doc, el, "ECFDF5", "059669")
            continue
        if "keybox" in classes:
            shade_block(doc, el, "F8FAFC", "64748B")
            continue
        if "summary-box" in classes:
            shade_block(doc, el, "F0FDFA", "0F766E")
            continue
        if "case" in classes:
            shade_block(doc, el, "F8FAFC", "0F766E")
            continue
        if "appendix-cover" in classes:
            add_page_break(doc)
            t1 = el.find(class_="t1")
            t2 = el.find(class_="t2")
            p = doc.add_heading(text_of(t1) + "  " + text_of(t2), level=1)
            for run in p.runs:
                set_run_font(run, size=16, bold=True, color=NAVY)
            rest = [c for c in el.find_all("div") if "t1" not in (c.get("class") or []) and "t2" not in (c.get("class") or [])]
            for c in rest:
                para = doc.add_paragraph()
                set_paragraph_spacing(para, after=8)
                r = para.add_run(text_of(c))
                set_run_font(r, size=11, color=TEAL)
            continue

    doc.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    convert()
