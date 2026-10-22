# -*- coding: utf-8 -*-
"""生成可印发的完整版纲要 Word 附件。"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "annex", "报告完整版纲要（已补全）.docx")


def set_run_font(run, size=12, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 0 else 14 if level == 1 else 12, bold=True, color=(4, 47, 46))
    return p


def add_p(doc, text, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    return p


def main():
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("旅顺口区海产品养殖加工产业专题分析报告")
    set_run_font(r, size=18, bold=True, color=(4, 47, 46))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("完整版纲要（已按正文补全）")
    set_run_font(r, size=14, bold=True, color=(15, 118, 110))
    add_p(doc, "编制日期：2026年10月22日。本纲要与 PDF 主报告、两菜专项附件、Excel 数据手册一一对应，不再保留空节。")

    add_heading(doc, "一、封面要点", 1)
    for t in [
        "底盘：2024年地方水产品产量25.7万吨；2023年海水养殖22万余吨、一产渔业总产值31亿元。",
        "两菜：海带3.4万亩、鲜品近30万吨；裙带菜6000余亩、鲜品约6万吨。",
        "种业：裙带菜苗市占约23%、海胆苗约30%，海带苗仅约2%。",
        "加工：110家企业、产值近20亿元；裙带菜加工收入5亿元、出口3.6亿元。",
        "路径：不盲目扩面，推进种业、提取、地标溢价和八项工程。",
    ]:
        add_p(doc, "• " + t)

    add_heading(doc, "二、必须并列的口径", 1)
    add_p(doc, "海带面积：生产口径3.4万亩（2025—2026年区海洋局采收季）；地标宣传口径6万余亩。产值：2026年采收季预计近4亿元；地标运用报道约30亿元。鲜品、烫煮后、干品、加工成品不得加总。")

    add_heading(doc, "三、正文十五章", 1)
    chapters = [
        ("一、研究背景、对象与方法", "完整版目的、边界、口径分层与工程化输出。"),
        ("二、战略定位与宏观环境", "大食物观、大连海洋经济、三类需求、近岸约束。"),
        ("三、资源禀赋与空间布局", "北纬39°、渔业用海39平方公里、柏岚子—龙头格局、生产年历。"),
        ("四、产业总体规模与结构", "公报指标、两菜占产量七成、110家加工企业。"),
        ("五、全产业链", "种苗、养殖、加工三层、流通、地标、金融工具箱。"),
        ("六、海带产业专题", "3.4万亩与2026行情、品质溢价、中宝1号、提取缺口。"),
        ("七、裙带菜产业专题", "全国份额、海宝1号/2号、盖世标准路径、出口。"),
        ("八、海珍品与其他品种", "海胆种业、海参冻干、鲍鱼地标、IMTA。"),
        ("九、市场、消费与国际贸易", "川渝火锅、国内预制菜、日美欧、价格指数。"),
        ("十、科技支撑、种业与绿色发展", "院所合作、机械化、废水、蓝碳。"),
        ("十一、对标分析", "连江规模、荣成提取、日韩分级。"),
        ("十二、问题诊断与风险", "P1—P7与气候、出口、环保三类风险。"),
        ("十三、发展思路、目标与路径", "有限海面无限价值链；2027/2030建议目标。"),
        ("十四、重点工程与项目库", "八项工程及与现有贷款、补贴的衔接。"),
        ("十五、保障措施与90日行动清单", "专班、要素和六件立刻可做的事。"),
    ]
    for name, desc in chapters:
        add_p(doc, name + "：" + desc)

    add_heading(doc, "四、附件", 1)
    add_p(doc, "附件A：两菜产业发展报告独立报送稿。")
    add_p(doc, "附件B：数据口径、企业名录、政策清单索引。")
    add_p(doc, "Excel《旅顺口区海产品产业数据手册》：九张工作表。")
    add_p(doc, "图1—图8：规模、种业、价值链、全国格局、主体、对标、结构、年历。")

    add_heading(doc, "五、原三份材料对应关系", 1)
    add_p(doc, "纲要 → 本文件 + 正文十五章全部展开。")
    add_p(doc, "海产品养殖加工产业专题分析报告 → 正文第一至五章、第八至十五章。")
    add_p(doc, "两菜产业发展报告 → 正文第六、七章 + 附件A。")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
