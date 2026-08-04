# -*- coding: utf-8 -*-
"""
丹东城市旅游价值提升策略报告 PPT 生成脚本
- 全部图表为 PowerPoint 原生图表(数据可在 PPT 内直接编辑)
- 全部示意图为原生形状与文本框(可自由编辑颜色/文字/位置)
运行: python3 build_dandong_ppt.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

# ---------------- 设计系统 ----------------
PRIMARY = RGBColor(0x14, 0x48, 0x7F)   # 鸭绿江深蓝
PRIMARY_D = RGBColor(0x0D, 0x2E, 0x52) # 更深蓝
ACCENT = RGBColor(0xE8, 0xA3, 0x3D)    # 银杏金
RED = RGBColor(0xB8, 0x3A, 0x2E)       # 红色文化
TEAL = RGBColor(0x2E, 0x8B, 0x83)      # 生态绿
DARK = RGBColor(0x22, 0x2B, 0x36)
GRAY = RGBColor(0x5E, 0x6B, 0x7A)
LIGHT = RGBColor(0xF0, 0xF4, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINEGRAY = RGBColor(0xC9, 0xD3, 0xDE)

FONT = "微软雅黑"
SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

PAGE_NO = [0]


def new_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, shadow=False):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, runs, size=14, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             space_after=0):
    """runs: str 或 [(text, dict_overrides), ...] 的段落列表"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    first = True
    for text, ov in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ov.get("align", align)
        p.line_spacing = ov.get("line_spacing", line_spacing)
        p.space_after = Pt(ov.get("space_after", space_after))
        segs = text if isinstance(text, list) else [(text, {})]
        for seg_text, seg_ov in segs:
            r = p.add_run()
            r.text = seg_text
            f = r.font
            f.name = FONT
            f.size = Pt(seg_ov.get("size", ov.get("size", size)))
            f.bold = seg_ov.get("bold", ov.get("bold", bold))
            f.color.rgb = seg_ov.get("color", ov.get("color", color))
    return tb


def header(slide, part, title, subtitle=None):
    """内容页顶部标题区"""
    PAGE_NO[0] += 1
    add_rect(slide, 0, 0, SW, Inches(0.12), ACCENT)
    add_rect(slide, Inches(0.55), Inches(0.42), Inches(0.09), Inches(0.62), PRIMARY)
    add_text(slide, Inches(0.85), Inches(0.30), Inches(10.5), Inches(0.4),
             part, size=11, color=ACCENT, bold=True)
    add_text(slide, Inches(0.85), Inches(0.58), Inches(11.5), Inches(0.55),
             title, size=24, color=PRIMARY_D, bold=True)
    if subtitle:
        add_text(slide, Inches(0.85), Inches(1.12), Inches(11.8), Inches(0.35),
                 subtitle, size=12, color=GRAY)
    # 页脚
    add_text(slide, Inches(0.55), Inches(7.12), Inches(6), Inches(0.3),
             "丹东城市旅游价值提升策略报告", size=9, color=LINEGRAY)
    add_text(slide, Inches(12.3), Inches(7.12), Inches(0.6), Inches(0.3),
             str(PAGE_NO[0] + 1), size=9, color=LINEGRAY, align=PP_ALIGN.RIGHT)


def kpi_card(slide, x, y, w, h, value, label, note=None, color=PRIMARY):
    add_rect(slide, x, y, w, h, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(slide, x, y, Inches(0.07), h, color)
    add_text(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.35), Inches(0.55),
             value, size=24, color=color, bold=True)
    add_text(slide, x + Inches(0.2), y + Inches(0.66), w - Inches(0.35), Inches(0.35),
             label, size=11.5, color=DARK, bold=True)
    if note:
        add_text(slide, x + Inches(0.2), y + Inches(0.98), w - Inches(0.35), h - Inches(1.05),
                 note, size=9.5, color=GRAY, line_spacing=1.05)


def bullets(slide, x, y, w, h, items, size=13, gap=6, line_spacing=1.12):
    """items: [(lead, body)] lead 加粗上色, body 常规"""
    paras = []
    for lead, body in items:
        segs = []
        if lead:
            segs.append((lead, {"bold": True, "color": PRIMARY}))
        segs.append((body, {}))
        paras.append(([("▎", {"color": ACCENT, "bold": True})] + segs
                      if False else segs, {"space_after": gap, "line_spacing": line_spacing}))
    # 加圆点前缀
    paras = []
    for lead, body in items:
        segs = [("● ", {"color": ACCENT, "size": size - 3})]
        if lead:
            segs.append((lead, {"bold": True, "color": PRIMARY_D}))
        segs.append((body, {}))
        paras.append((segs, {"space_after": gap, "line_spacing": line_spacing}))
    add_text(slide, x, y, w, h, paras, size=size, color=DARK)


def style_chart(chart, font_size=10, legend=True, data_labels=True,
                num_fmt='0.0"%"', label_size=9):
    chart.font.name = FONT
    chart.font.size = Pt(font_size)
    chart.font.color.rgb = DARK
    if legend:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font_size)
    else:
        chart.has_legend = False
    if data_labels:
        for plot in chart.plots:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(label_size)
            dl.font.bold = True
            dl.font.color.rgb = PRIMARY_D
            dl.number_format = num_fmt
            dl.number_format_is_linked = False
            try:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            except Exception:
                pass


def add_col_chart(slide, x, y, w, h, cats, series, title=None,
                  colors=None, num_fmt='0.0"%"', legend=None,
                  chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED, label_size=9):
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    gframe = slide.shapes.add_chart(chart_type, x, y, w, h, data)
    chart = gframe.chart
    if legend is None:
        legend = len(series) > 1
    style_chart(chart, legend=legend, num_fmt=num_fmt, label_size=label_size)
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        tp = chart.chart_title.text_frame.paragraphs[0]
        tp.font.size = Pt(12)
        tp.font.bold = True
        tp.font.name = FONT
        tp.font.color.rgb = PRIMARY_D
    else:
        chart.has_title = False
    palette = colors or [PRIMARY, ACCENT, TEAL, RED]
    for i, s in enumerate(chart.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = palette[i % len(palette)]
    try:
        chart.plots[0].gap_width = 60
        chart.plots[0].overlap = -20 if len(series) > 1 else 0
    except Exception:
        pass
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE4, 0xEA, 0xF1)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(9)
    va.tick_labels.font.name = FONT
    ca = chart.category_axis
    ca.format.line.color.rgb = LINEGRAY
    ca.tick_labels.font.size = Pt(9.5)
    ca.tick_labels.font.name = FONT
    return chart


def add_pie_chart(slide, x, y, w, h, cats, vals, title=None, colors=None):
    data = CategoryChartData()
    data.categories = cats
    data.add_series("占比", vals)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.PIE, x, y, w, h, data)
    chart = gframe.chart
    chart.font.name = FONT
    chart.font.size = Pt(10)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0"%"'
    dl.number_format_is_linked = False
    dl.font.size = Pt(10)
    dl.font.bold = True
    dl.font.color.rgb = WHITE
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        tp = chart.chart_title.text_frame.paragraphs[0]
        tp.font.size = Pt(12); tp.font.bold = True
        tp.font.name = FONT; tp.font.color.rgb = PRIMARY_D
    else:
        chart.has_title = False
    palette = colors or [PRIMARY, ACCENT, TEAL, GRAY, RED]
    for i, pt in enumerate(chart.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = palette[i % len(palette)]
    return chart


def add_table(slide, x, y, w, h, rows, col_widths=None, header_fill=PRIMARY,
              font_size=10.5, header_size=11, row_h=None, align_first_left=True):
    n_r, n_c = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_r, n_c, x, y, w, h)
    tbl = shp.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(w * cw / total))
    for ri, row in enumerate(rows):
        if row_h:
            tbl.rows[ri].height = row_h
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if ri % 2 == 1 else LIGHT
            tf = cell.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (ci == 0 and align_first_left) else PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(val)
            f = r.font; f.name = FONT
            f.size = Pt(header_size if ri == 0 else font_size)
            f.bold = (ri == 0) or (ci == 0)
            f.color.rgb = WHITE if ri == 0 else (PRIMARY_D if ci == 0 else DARK)
    return tbl


def section_slide(no, title, points):
    s = new_slide()
    set_bg(s, PRIMARY_D)
    PAGE_NO[0] += 1
    add_rect(s, 0, Inches(6.9), SW, Inches(0.6), PRIMARY)
    add_rect(s, Inches(0.9), Inches(1.7), Inches(1.5), Inches(0.12), ACCENT)
    add_text(s, Inches(0.9), Inches(2.0), Inches(3), Inches(1.2),
             f"PART {no}", size=40, color=ACCENT, bold=True)
    add_text(s, Inches(0.9), Inches(3.0), Inches(11.4), Inches(1.0),
             title, size=34, color=WHITE, bold=True)
    add_text(s, Inches(0.95), Inches(4.15), Inches(10.8), Inches(2.2),
             [(f"—  {p}", {"space_after": 8}) for p in points],
             size=14.5, color=RGBColor(0xC7, 0xD6, 0xE8))
    add_text(s, Inches(12.15), Inches(7.0), Inches(0.8), Inches(0.35),
             str(PAGE_NO[0] + 1), size=10, color=RGBColor(0x8F, 0xA6, 0xC0),
             align=PP_ALIGN.RIGHT)
    return s


def note_bar(slide, x, y, w, text, color=ACCENT, h=Inches(0.52), size=11.5):
    add_rect(slide, x, y, w, h, RGBColor(0xFB, 0xF3, 0xE2) if color == ACCENT else LIGHT,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, x + Inches(0.18), y, w - Inches(0.36), h,
             [([("观点  ", {"bold": True, "color": RED}), (text, {})], {})],
             size=size, color=DARK, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================
# S1 封面
# ============================================================
s = new_slide()
set_bg(s, PRIMARY_D)
add_rect(s, 0, 0, SW, Inches(0.18), ACCENT)
add_rect(s, 0, Inches(7.32), SW, Inches(0.18), ACCENT)
add_rect(s, Inches(0.95), Inches(1.35), Inches(2.2), Inches(0.5), ACCENT,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.95), Inches(1.35), Inches(2.2), Inches(0.5),
         "策 略 研 究 报 告", size=14, color=PRIMARY_D, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.9), Inches(2.15), Inches(11.8), Inches(1.0),
         "从「双起点」到「留量之城」", size=42, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(3.15), Inches(11.8), Inches(0.7),
         "丹东城市旅游价值提升策略报告", size=26, color=ACCENT, bold=True)
add_text(s, Inches(0.93), Inches(4.15), Inches(11.5), Inches(0.9),
         [("人群变化与消费特征分析  ·  东北及全国城市定位比较  ·  IP打造 / 沿江品质 / 票务组合 / 产品体系", {}),
          ("银发经济 · 人口回流 · 城市更新的更高维度价值发掘", {})],
         size=14, color=RGBColor(0xC7, 0xD6, 0xE8), line_spacing=1.5)
# 底部三个标签
tags = ["G331 · G228 双国道零起点", "中国最大边境城市", "鸭绿江畔 · 英雄之城 · 银杏之城"]
tx = Inches(0.93)
for t in tags:
    wbox = Inches(3.6)
    add_rect(s, tx, Inches(5.6), wbox, Inches(0.55), None, line=RGBColor(0x5B, 0x79, 0x9E),
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, tx, Inches(5.6), wbox, Inches(0.55), t, size=12.5,
             color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tx += wbox + Inches(0.35)
add_text(s, Inches(0.93), Inches(6.65), Inches(11), Inches(0.4),
         "2026年7月  ·  提纲确认稿(V1)", size=12, color=RGBColor(0x8F, 0xA6, 0xC0))

# ============================================================
# S2 目录
# ============================================================
s = new_slide()
set_bg(s)
header(s, "CONTENTS", "目录", None)
toc = [
    ("01", "趋势与人群:谁在旅行,怎么消费", "旅游人群结构变化 · 银发/Z世代双极 · 消费特征与出行方式"),
    ("02", "定位比较:东北格局与全国样本", "东北城市旅游定位图谱 · 延吉对标 · 河南文旅与非顶流城市出圈路径"),
    ("03", "核心观点:丹东的八个判断", "缝隙定位 · 从过路到留量 · 银发主场 · 起点仪式经济"),
    ("04", "策略主张:IP · 沿江 · 联动 · 票务 · 产品", "双起点之城 · 虎山—断桥黄金岸线 · 一江两线三圈 · City Pass"),
    ("05", "客群导入与更高维度价值", "客群分层与转化 · 银发经济 · 人口回流 · 城市更新"),
    ("06", "实施路径与保障", "三阶段路线图 · 运营机制 · 风险边界"),
]
y = Inches(1.65)
for no, t, d in toc:
    add_rect(s, Inches(0.9), y, Inches(11.5), Inches(0.82), LIGHT,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(1.15), y + Inches(0.1), Inches(0.9), Inches(0.6),
             no, size=24, color=ACCENT, bold=True)
    add_text(s, Inches(2.15), y + Inches(0.09), Inches(5.6), Inches(0.4),
             t, size=15.5, color=PRIMARY_D, bold=True)
    add_text(s, Inches(2.15), y + Inches(0.47), Inches(9.9), Inches(0.32),
             d, size=10.5, color=GRAY)
    y += Inches(0.94)

# ============================================================
# PART 01  趋势与人群
# ============================================================
section_slide("01", "趋势与人群:谁在旅行,怎么消费",
              ["旅游市场从规模扩张转入结构重构:客群、渠道、消费逻辑同时生变",
               "银发与Z世代形成「双极增长」,高铁+自驾重塑出行半径",
               "体验溢价取代门票经济,县域与边境成为新的价值洼地"])

# S: 五大结构性变化
s = new_slide(); set_bg(s)
header(s, "PART 01 · 趋势与人群", "旅游消费的五大结构性变化",
       "2025—2026年市场数据显示:需求没有减弱,而是被重新分配")
cards = [
    ("观光退潮\n体验为王", "传统门票收入下滑,非遗/沉浸/研学订单暴涨40%–80%;游客愿为独特体验多付30%–100%", PRIMARY),
    ("双极增长\n代际分化", "Z世代占文旅消费人群约40%;银发线上订单同比+20%以上,61–65岁增速超50%", ACCENT),
    ("高铁+自驾\n重塑半径", "高铁取代飞机成中长途首选(51% vs 42%);节假日租车自驾预订同比+113%", TEAL),
    ("县域替代\n反向旅游", "县域及以下客源占比超30%,市场规模约1.2万亿;本质是「体验升级的空间替代」", RED),
    ("情绪价值\n圈层传播", "决策依赖社交平台与圈层口碑;悦己、治愈、陪伴、仪式感成为核心购买理由", PRIMARY_D),
]
x = Inches(0.62)
for t, d, c in cards:
    w = Inches(2.34)
    add_rect(s, x, Inches(1.75), w, Inches(3.3), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.75), w, Inches(0.12), c)
    add_text(s, x + Inches(0.16), Inches(2.05), w - Inches(0.32), Inches(0.95),
             t, size=15, color=c, bold=True, line_spacing=1.1)
    add_text(s, x + Inches(0.16), Inches(3.05), w - Inches(0.32), Inches(1.9),
             d, size=10.5, color=DARK, line_spacing=1.2)
    x += w + Inches(0.12)
note_bar(s, Inches(0.62), Inches(5.45), Inches(12.1),
         "需求端的每一项变化,都在降低「资源型头部景区」的权重、抬高「体验型内容城市」的权重——这正是丹东这类非顶流城市的历史性窗口。")
add_text(s, Inches(0.62), Inches(6.25), Inches(12.1), Inches(0.7),
         "数据来源:新旅界《2021–2025中国文旅市场消费洞察》、2026年五一假期旅游数据全景报告、2026Q2中国旅游消费趋势调研(n=3000)等公开资料",
         size=9, color=GRAY)

# S: 人群结构 双极增长
s = new_slide(); set_bg(s)
header(s, "PART 01 · 趋势与人群", "人群结构:银发与Z世代的「双极增长」",
       "客群基本盘从「中间大、两头小」转为「两头强」——供给必须精准分层")
add_col_chart(s, Inches(0.6), Inches(1.7), Inches(6.3), Inches(4.3),
              ["Z世代占文旅\n消费人群", "亲子家庭\n出游占比", "50岁+境内游\n出行比例", "县域及以下\n客源占比"],
              [("占比(%)", (40, 40, 41.8, 30))],
              title="主要客群关键占比指标(2025–2026,多源口径)",
              colors=[PRIMARY], num_fmt='0.0"%"', legend=False, label_size=10)
bullets(s, Inches(7.25), Inches(1.8), Inches(5.5), Inches(3.6), [
    ("Z世代(18–28岁):", "总量约2.3亿并已见顶,依赖流量红利的打法失效,必须转向提升人均消费价值;决策看情绪、颜值、社交属性。"),
    ("银发族(55岁+):", "2025年60岁+人口占比超23%,银发文旅年复合增长12%以上,是确定性最强的赛道;61–65岁订单增速超50%。"),
    ("亲子家庭:", "占出游约40%,研学、科普、自然类内容是刚需;三代同游比例上升,对住宿舒适度与安全性要求高。"),
    ("单人出行:", "户均人口降至2.3、单人家庭超25%,一人游/宠物友好等新需求爆发。"),
], size=12, gap=8)
note_bar(s, Inches(7.25), Inches(5.55), Inches(5.5),
         "丹东客群策略必须「双线作战」:用仪式感与内容抓年轻圈层,用康养旅居抓银发基本盘。", size=10.5, h=Inches(0.75))

# S: 银发深描
s = new_slide(); set_bg(s)
header(s, "PART 01 · 趋势与人群", "银发客群深描:被低估的高价值增量",
       "「低频、低消、低关注」的旧认知正在被数据颠覆")
add_col_chart(s, Inches(0.6), Inches(1.7), Inches(6.1), Inches(4.2),
              ["50岁+线上订单\n同比增速", "61–65岁群体\n订单增速", "银发文旅市场\n年复合增长", "银发人均住宿花费\n高于青年群体"],
              [("增速/差值(%)", (20, 50, 12, 30))],
              title="银发文旅关键增长指标(2025)",
              colors=[ACCENT], num_fmt='0"%"', legend=False, label_size=10)
bullets(s, Inches(7.05), Inches(1.75), Inches(5.7), Inches(3.5), [
    ("时间自由 + 消费稳定:", "可避开高峰形成「淡季不淡」的平峰需求,退休金收入抗周期性强。"),
    ("内部显著分化:", "50–60岁「准银发」偏好文化研学,行为接近新中产;60–75岁「活力银发」是康养旅居主力;75岁+重医疗与安全。"),
    ("高端消费主力:", "豪华邮轮客群中50–70岁占比超70%,定制游单次花费1.2万元以上,较其他年龄组高约25%。"),
    ("季节迁徙规律:", "夏避暑、冬避寒、春赏花,目的地相对固定——一旦建立心智即形成稳定复游。"),
], size=12, gap=8)
note_bar(s, Inches(7.05), Inches(5.5), Inches(5.7),
         "银发要的不是低价,而是「高品质、慢节奏、有保障」——丹东江海泉组合+低生活成本,恰好匹配。", size=10.5, h=Inches(0.75))

# S: 消费特征
s = new_slide(); set_bg(s)
header(s, "PART 01 · 趋势与人群", "消费特征:客单价分层与「体验溢价」",
       "人均消费横盘,但花钱的地方彻底变了——从门票转向内容、住宿与餐饮")
add_col_chart(s, Inches(0.6), Inches(1.7), Inches(5.9), Inches(4.2),
              ["全国国内游\n人均消费", "一二线城市\n旅游客单价", "县域旅游\n客单价"],
              [("元/人次", (961, 1200, 850))],
              title="旅游客单价分层(2026,元)",
              colors=[PRIMARY], num_fmt='0"元"', legend=False, label_size=10)
bullets(s, Inches(6.85), Inches(1.75), Inches(5.9), Inches(3.4), [
    ("门票经济退潮:", "传统景点门票收入下滑,游客对门票价格高度敏感,对「一票多日」「城市通票」接受度高。"),
    ("体验溢价崛起:", "为独一无二、有记忆点的体验多付30%–100%;非遗体验类订单同比+40%–80%。"),
    ("二消结构升级:", "特色住宿、在地美食、文创与演艺成为消费主力项;收入结构应「票务让利、二消挣钱」。"),
    ("停留时间=收入:", "县域黑马目的地停留从1.5天拉长至3天以上——延长停留是提升总消费的第一杠杆。"),
], size=12, gap=8)
note_bar(s, Inches(6.85), Inches(5.4), Inches(5.9),
         "丹东不应追求高票价,而应以「低门槛进入 + 高体验二消 + 长停留」重构收入模型。", size=10.5, h=Inches(0.75))

# S: 出行方式
s = new_slide(); set_bg(s)
header(s, "PART 01 · 趋势与人群", "出行方式:高铁+自驾的「新半径」",
       "交通方式变化直接决定客源腹地——这是丹东最被低估的变量")
add_pie_chart(s, Inches(0.7), Inches(1.75), Inches(5.2), Inches(4.4),
              ["高铁", "飞机", "其他(自驾/大巴等)"], [51, 42, 7],
              title="中长途出行首选交通方式(2026)")
bullets(s, Inches(6.35), Inches(1.75), Inches(6.4), Inches(3.3), [
    ("高铁重构客源圈:", "沈白高铁(2025年9月)通车、丹大快铁成熟,沈阳/大连2小时圈、京津冀5小时圈成为丹东现实腹地。"),
    ("自驾成为增长极:", "节假日租车预订同比+113%;「高铁+落地自驾」成为主流组合,自驾IP线路(318/G219/G331)热度持续攀升。"),
    ("公路本身即目的地:", "G331被纳入国家「黄金大外环」,「此生必驾」的仪式性需求兴起——起点城市天然获得首站流量。"),
    ("房车与营地爆发:", "丹东国门湾营地日接待已近300辆房车,车位一位难求——供给远未跟上需求。"),
], size=12, gap=8)
note_bar(s, Inches(6.35), Inches(5.35), Inches(6.4),
         "对丹东,自驾客群不是「过路客」问题的根源,而是最大的增量入口——关键在于用产品把他们留下来。", size=10.5, h=Inches(0.75))

# S: 小结 对丹东的含义
s = new_slide(); set_bg(s)
header(s, "PART 01 · 趋势与人群", "小结:人群与消费变化对丹东的六个含义", None)
imps = [
    ("① 窗口期真实存在", "反向旅游与边境游升温,使非顶流城市第一次有机会靠「内容+服务」而非「资源等级」赢得客源。"),
    ("② 自驾是第一入口", "G331/G228双起点叠加自驾大盘增长,丹东获得了一个不需要花钱买的全国性流量入口。"),
    ("③ 银发是第二曲线", "东北夏季凉爽+温泉+低生活成本,与银发「避暑旅居」需求高度匹配,可对冲淡季。"),
    ("④ 门票逻辑必须放弃", "客群对门票敏感、对体验慷慨——票务组合要做「进入门槛最低化」,收入靠停留与二消。"),
    ("⑤ 内容决定溢价", "红色叙事、边境风情、江海景观都需转译为可参与、可传播的体验内容,才能兑现30%–100%溢价。"),
    ("⑥ 圈层渠道>大众广告", "自驾俱乐部、康养机构、研学机构等B端圈层渠道的转化效率远高于泛投放。"),
]
x, y = Inches(0.62), Inches(1.65)
for i, (t, d) in enumerate(imps):
    cx = x + (i % 3) * Inches(4.18)
    cy = y + (i // 3) * Inches(2.45)
    add_rect(s, cx, cy, Inches(4.0), Inches(2.28), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, cx, cy, Inches(0.08), Inches(2.28), ACCENT)
    add_text(s, cx + Inches(0.22), cy + Inches(0.15), Inches(3.65), Inches(0.45),
             t, size=14.5, color=PRIMARY_D, bold=True)
    add_text(s, cx + Inches(0.22), cy + Inches(0.62), Inches(3.65), Inches(1.6),
             d, size=11, color=DARK, line_spacing=1.2)

# ============================================================
# PART 02  定位比较
# ============================================================
section_slide("02", "定位比较:东北格局与全国样本",
              ["东北旅游呈「哈尔滨超级IP+沈大枢纽+长白山生态」格局,丹东是规划中的集散/特色城市",
               "延吉证明:边境+民族风情的中等城市可以做成全国性目的地",
               "河南与淄博/天水等样本给出非资源顶流城市的出圈方法论"])

# S: 东北定位图谱
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "东北旅游城市定位图谱",
       "《东北地区旅游业发展规划》与市场表现共同定义的分工格局")
rows = [
    ["城市", "规划/市场定位", "核心IP与心智", "主力客群", "对丹东的启示"],
    ["哈尔滨", "世界级冰雪旅游城市", "冰雪大世界 · 中央大街 · 「尔滨」宠客", "全国年轻客群+入境", "服务口碑可再造城市;IP需持续宠客运营"],
    ["沈阳", "东北亚旅游消费中心", "故宫中街 · 文体演艺 · 洗浴文化", "都市休闲+过夜消费", "「消费场景城市」路线,非资源依赖"],
    ["大连", "滨海度假中心城市", "浪漫海滨 · 邮轮 · 避暑", "华北华东度假客", "滨海度假与丹东江海资源部分同质,需错位"],
    ["长白山/延边", "生态冰雪目的地", "长白山天池 · 延吉民族风情", "自驾+高铁圈层客群", "边境+民族风情可成全国性IP(延吉样本)"],
    ["丹东", "重点旅游集散城市 · 特色旅游城市", "双国道起点 · 断桥 · 英雄城市 · 银杏", "待建立稳定心智", "规划身份≠市场心智,需自建尖锐记忆点"],
]
add_table(s, Inches(0.55), Inches(1.68), Inches(12.25), Inches(4.15), rows,
          col_widths=[1.0, 1.6, 2.1, 1.5, 2.3], font_size=10.5, header_size=11.5)
note_bar(s, Inches(0.55), Inches(6.1), Inches(12.25),
         "东北旅游的头部格局已经封盘(冰雪=哈尔滨、滨海=大连、消费=沈阳),丹东不应正面竞争,而应占据尚无人占据的「边境公路起点+英雄叙事+旅居暖城」缝隙。")

# S: 东北数据比较
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "东北重点城市旅游数据比较(2025)",
       "丹东体量差距大,但增速处于第一梯队——弹性大、基数低是后发优势")
add_col_chart(s, Inches(0.6), Inches(1.7), Inches(6.4), Inches(4.2),
              ["哈尔滨", "大连", "丹东", "辽宁全省\n(2024)"],
              [("游客人次增速(%)", (12.7, 24.0, 19.3, 28.1)),
               ("旅游收入增速(%)", (21.8, 24.0, 23.8, 25.9))],
              title="2025年旅游接待与收入同比增速比较",
              colors=[PRIMARY, ACCENT], num_fmt='0.0"%"')
rows = [
    ["城市/指标", "接待规模(2025)", "标志性事实"],
    ["哈尔滨", "2.02亿人次 / 2818亿元", "联合国「世界冰雪旅游卓越城市」;入境+57.7%"],
    ["延边(延吉)", "国庆543万人次 / 34.4亿元", "边境民族风情出圈;沈白高铁直达"],
    ["沈阳", "国庆1560万人次 / 141亿元", "文体旅消费场景之城"],
    ["丹东", "全年人次+19.3% / 收入+23.8%", "断桥元旦+102%;国门湾营地日均近300辆房车"],
]
add_table(s, Inches(7.25), Inches(1.75), Inches(5.5), Inches(3.6), rows,
          col_widths=[1.0, 1.5, 2.2], font_size=9.5, header_size=10.5)
note_bar(s, Inches(7.25), Inches(5.6), Inches(5.5),
         "丹东与哈尔滨的差距不是「差一个量级的营销」,而是差一套「让人留下来」的产品与内容体系。", size=10.5, h=Inches(0.75))

# S: 延吉对标
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "延吉样本:与丹东最可比的边境城市",
       "同为边境、同靠民族/异域风情、同非传统旅游强市——延吉先走通了这条路")
rows = [
    ["维度", "延吉(延边州)", "丹东", "差距/机会点"],
    ["城市心智", "「小首尔」异域风情,弹幕墙/民族服饰写真全网刷屏", "断桥+英雄城市认知存量高,但缺少年轻化转译", "丹东认知资产更厚,缺的是「可拍、可穿、可晒」的体验化表达"],
    ["体验产品", "民族服饰写真+水上市场+特色餐饮,人均消费场景密集", "游船、景区观光为主,换装/写真/市集等参与型产品稀缺", "把「朝鲜族风情+江景」做成沉浸式消费场景"],
    ["交通引流", "沈白高铁开通即承接长白山客流,「一程多站」枢纽", "双国道零起点+丹大快铁,自驾流量自发聚集", "把自驾流量的「首站仪式」固化为标准产品"],
    ["流量运营", "政府+商户高频造节,社媒话题持续更新", "「丹东真好」口号已立,话题工程刚起步", "建立常态化内容生产机制,而非节庆式脉冲"],
]
add_table(s, Inches(0.55), Inches(1.68), Inches(12.25), Inches(4.3), rows,
          col_widths=[0.75, 1.9, 1.9, 2.0], font_size=10, header_size=11)
note_bar(s, Inches(0.55), Inches(6.2), Inches(12.25),
         "延吉证明「边境风情是可以卖给全国年轻人的」;丹东的资源组合(江+海+桥+长城+温泉)比延吉更立体,差在产品化与运营密度。")

# S: 全国非顶流城市出圈比较
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "全国样本:资源非顶流城市的出圈路径",
       "五个样本,五种引爆点,一个共同规律")
rows = [
    ["城市", "引爆点", "承接与转化", "长红关键", "可迁移到丹东"],
    ["淄博", "烧烤+大学生「进淄赶烤」", "政府兜底服务(专列/公交/管控物价)", "真诚服务口碑", "「不让游客受委屈」的服务承诺与先行赔付"],
    ["天水", "麻辣烫单品引爆", "美食街区快速扩容", "单品→全域动线", "黄蚬子/开海节做成全国性美食事件"],
    ["榕江", "村超全民赛事", "赛事日历化+农产品带货", "全民参与感", "鸭绿江马拉松/自驾集结赛事IP化"],
    ["泉州", "簪花围民俗视觉符号", "写真产业链+古城漫游", "视觉符号自传播", "银杏大道/朝鲜族服饰的「可穿戴符号」"],
    ["开封(万岁山)", "武侠NPC+百元票三日无限次", "高频内容更新+玩家化运营", "极致性价比+沉浸", "通票制+常态化NPC演艺(安东老街/游船)"],
]
add_table(s, Inches(0.55), Inches(1.65), Inches(12.25), Inches(4.35), rows,
          col_widths=[0.95, 1.5, 1.6, 1.2, 2.1], font_size=10, header_size=11)
note_bar(s, Inches(0.55), Inches(6.2), Inches(12.25),
         "共同规律:尖锐记忆点引爆 → 政府服务兜底口碑 → 产品快速承接流量 → 事件日历化维持热度。丹东的「尖锐记忆点」应锁定:此生必驾,始于丹东。")

# S: 河南方法论
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "河南文旅方法论:从「资源」到「内容」的完整证明",
       "河南用人造内容跑赢了资源禀赋——这对丹东是最重要的战略参照")
mcards = [
    ("万岁山武侠城", "低价高频 + 玩家化运营", "百元门票三日无限次;NPC全域互动;2025年前三季度1550万人次、收入9.5亿元,同比+170%/+167%", ACCENT),
    ("只有河南·戏剧幻城", "在地文化的沉浸式精品", "21个剧场、约700分钟剧目;开城至今观剧超5000万人次,85后占85%、八成来自省外", PRIMARY),
    ("洛邑古城/洛阳", "可穿戴的文化符号", "汉服写真产业链让游客成为内容生产者;大遗址保护与城市更新深度融合", TEAL),
    ("清明上河园/云台山", "双IP长线 + 精细服务", "名画IP长期运营;「不让游客受委屈」服务体系入选全国案例", RED),
]
x = Inches(0.62)
for t, sub, d, c in mcards:
    w = Inches(2.98)
    add_rect(s, x, Inches(1.7), w, Inches(3.35), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.7), w, Inches(0.12), c)
    add_text(s, x + Inches(0.18), Inches(1.98), w - Inches(0.36), Inches(0.42),
             t, size=13.5, color=c, bold=True)
    add_text(s, x + Inches(0.18), Inches(2.42), w - Inches(0.36), Inches(0.4),
             sub, size=11, color=PRIMARY_D, bold=True)
    add_text(s, x + Inches(0.18), Inches(2.86), w - Inches(0.36), Inches(2.1),
             d, size=10, color=DARK, line_spacing=1.22)
    x += w + Inches(0.12)
add_text(s, Inches(0.62), Inches(5.3), Inches(12.1), Inches(0.4),
         "方法论提炼(可迁移,不照搬形态):", size=13, color=PRIMARY_D, bold=True)
add_text(s, Inches(0.62), Inches(5.72), Inches(12.1), Inches(1.1),
         [([("产品为王", {"bold": True, "color": RED}), ("(内容即吸引物,而非资源等级)  ·  ", {}),
            ("沉浸互动", {"bold": True, "color": RED}), ("(游客从旁观者变参与者)  ·  ", {}),
            ("极致性价比", {"bold": True, "color": RED}), ("(低门槛进入换取长停留与二消)  ·  ", {}),
            ("高频更新", {"bold": True, "color": RED}), ("(内容日历化对抗审美疲劳)  ·  ", {}),
            ("真诚服务", {"bold": True, "color": RED}), ("(口碑是最便宜的获客渠道)", {})], {"line_spacing": 1.4})],
         size=12.5, color=DARK)

# S: 丹东现状与短板
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "丹东现状:动能已现,短板清晰",
       "流量正在自发到来,但供给侧还接不住")
add_col_chart(s, Inches(0.6), Inches(1.7), Inches(6.2), Inches(4.15),
              ["断桥景区\n(2026元旦)", "鸭绿江一号码头\n(2026元旦)", "全市游客量\n(2026元旦)", "全市旅游收入\n(2026元旦)"],
              [("同比增速(%)", (102, 350, 22.73, 33.62))],
              title="丹东旅游动能指标(同比增速)",
              colors=[TEAL], num_fmt='0.0"%"', legend=False, label_size=10)
bullets(s, Inches(7.05), Inches(1.72), Inches(5.7), Inches(4.3), [
    ("动能:", "2025年全市旅游人数+19.3%、收入+23.8%;G331「零公里」出圈,国门湾营地一位难求;「鸭绿江畔·丹东真好」品牌起势。"),
    ("短板① 无龙头:", "全市30处A级景区、9个4A,没有5A级景区,缺乏可以「扛心智」的旗舰产品。"),
    ("短板② 过路化:", "自驾客打卡零公里即北上,停留短、消费浅;夜间产品与沉浸内容近乎空白。"),
    ("短板③ 断点多:", "虎山—断桥沿江体验不连续,景区散而不联,一票一景、各自为战。"),
    ("短板④ 季节失衡:", "旺季集中于夏秋,冬春供给单薄;住宿结构老化,中高端与特色民宿稀缺。"),
], size=11.5, gap=6)

# S: 定位结论
s = new_slide(); set_bg(s)
header(s, "PART 02 · 定位比较", "比较结论:丹东的定位选择",
       "不做「东北的哈尔滨」,做「中国边境旅行的起点与暖城」")
add_rect(s, Inches(0.62), Inches(1.75), Inches(12.1), Inches(1.5), PRIMARY_D,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(1.0), Inches(1.95), Inches(11.4), Inches(1.1),
         [([("城市定位主张:", {"color": ACCENT, "bold": True, "size": 15}),
            ("  中国边境旅行第一站 · 英雄暖城丹东", {"color": WHITE, "bold": True, "size": 20})], {}),
          ("以「双起点」占据全国自驾/公路旅行心智,以「江海泉+低成本」占据银发旅居心智,以「英雄叙事」守住红色研学基本盘。",
           {"color": RGBColor(0xC7, 0xD6, 0xE8), "size": 12.5, "line_spacing": 1.3})],
         size=14)
pos = [
    ("对哈尔滨", "不争冰雪超级IP,承接其溢出客流:「冰雪之后,来暖城过冬/沿边出发」;加入东北高铁计次票等联动产品。"),
    ("对大连/沈阳", "不争滨海度假与都市消费,做「边境体验+公路旅行」差异内容,吸引其2小时圈周末客群。"),
    ("对延吉", "共同做大「边境风情」品类,以「起点仪式+英雄叙事+江海景观」形成差异;争取G331沿线城市联盟主导权。"),
    ("对全国", "锁定两个全国唯一:双国道零起点、最大边境城市——所有传播资源向「此生必驾,始于丹东」集中。"),
]
y = Inches(3.55)
for t, d in pos:
    add_rect(s, Inches(0.62), y, Inches(12.1), Inches(0.74), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(0.85), y, Inches(1.7), Inches(0.74), t, size=12.5,
             color=PRIMARY_D, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.7), y, Inches(9.9), Inches(0.74), d, size=11.5,
             color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.85)

# ============================================================
# PART 03  核心观点
# ============================================================
section_slide("03", "核心观点:关于丹东的八个判断",
              ["基于人群变化、消费特征与城市比较得出的策略判断",
               "每个观点对应后文一组可落地的策略动作"])

views1 = [
    ("观点一|丹东的对手不是别的城市,而是「过路化」",
     "双起点带来的自驾流量是自发的、免费的,但目前多数客人「打卡零公里即离开」。把平均停留从半天拉到2晚,总消费的提升幅度将远大于游客数翻倍,这是全部策略的第一目标。"),
    ("观点二|「起点」是仪式经济,不是地理概念",
     "游客要的是「万里征程始于足下」的仪式感:发车礼、第一枚章、纪念车贴、护照式打卡。把仪式做成标准化产品,零公里地标才能从拍照点变成消费入口。"),
    ("观点三|红色叙事是护城河,但必须庄重地年轻化",
     "抗美援朝记忆是丹东独有的、不可复制的心智资产。用沉浸式讲述(1950实景演艺、深度研学课程)替代静态陈列,但娱乐化必须有边界——庄重本身就是体验的一部分。"),
    ("观点四|银发旅居是丹东的「第二人口」策略",
     "东北夏凉、温泉、江海、低房价低物价——银发避暑旅居可以一次性解决淡季现金流、存量房产消化和服务业就业三个问题,应升格为城市战略而非文旅子项。"),
]
views2 = [
    ("观点五|票务上做「减法」,收入上做「加法」",
     "客群对门票敏感、对体验慷慨。学万岁山:通票化、多日多次、极限淡季定价,把收入重心移向住宿、餐饮、演艺、文创——先让人进来、留下,钱自然花出来。"),
    ("观点六|沿江20公里是丹东的「中央体验区」",
     "虎山—断桥黄金岸线承担了游客80%的第一印象,应按「一带四节点」整体运营:仪式区(虎山)—集散区(国门湾)—生活区(老城滨江)—核心区(断桥),夜航+夜市+夜演补齐夜间断档。"),
    ("观点七|联动的本质是「让游客的下一步永远有去处」",
     "一江两线三圈+城市通票+接驳专线,把散落的景区编织成连续动线;与G331沿线城市结盟,把「过境通道」变成「共同市场」。"),
    ("观点八|口碑是丹东唯一付得起的全国性广告",
     "淄博与哈尔滨证明:服务兜底产生的自来水传播,效果超过任何投放。「不让游客受委屈」应成为全市动员的一号工程。"),
]
for title_txt, views in [("核心观点(一):定位与叙事", views1), ("核心观点(二):模式与运营", views2)]:
    s = new_slide(); set_bg(s)
    header(s, "PART 03 · 核心观点", title_txt, None)
    y = Inches(1.72)
    for t, d in views:
        add_rect(s, Inches(0.62), y, Inches(12.1), Inches(1.22), LIGHT,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_rect(s, Inches(0.62), y, Inches(0.08), Inches(1.22), RED)
        add_text(s, Inches(0.9), y + Inches(0.1), Inches(11.6), Inches(0.4),
                 t, size=13.5, color=PRIMARY_D, bold=True)
        add_text(s, Inches(0.9), y + Inches(0.5), Inches(11.6), Inches(0.68),
                 d, size=11, color=DARK, line_spacing=1.18)
        y += Inches(1.33)

# ============================================================
# PART 04  策略主张
# ============================================================
section_slide("04", "策略主张:IP · 沿江 · 联动 · 票务 · 产品",
              ["城市IP:双起点之城的主副IP矩阵与仪式感系统",
               "沿江品质:虎山—断桥「一带四节点」中央体验区",
               "全域联动「一江两线三圈」与City Pass票务组合",
               "产品组合:自驾优先、四季日历、沉浸内容分期"])

# S: IP体系
s = new_slide(); set_bg(s)
header(s, "PART 04 · 策略主张", "城市IP体系:「此生必驾,始于丹东」",
       "一个主IP统领心智,四个副IP分季节、分客群展开")
add_rect(s, Inches(3.4), Inches(1.7), Inches(6.5), Inches(1.15), PRIMARY_D,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(3.4), Inches(1.82), Inches(6.5), Inches(0.5),
         "主IP:中国边境旅行第一站", size=17, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(3.4), Inches(2.32), Inches(6.5), Inches(0.45),
         "「此生必驾,始于丹东」— G331/G228 双国道零起点", size=12, color=WHITE, align=PP_ALIGN.CENTER)
subs = [
    ("英雄之城", "红色研学 · 抗美援朝纪念叙事", RED),
    ("银杏之城", "秋季流量引擎 · 城市视觉符号", ACCENT),
    ("江海泉暖城", "康养旅居 · 银发候鸟基地", TEAL),
    ("边境风情之城", "跨境视野 · 朝鲜族文化体验", PRIMARY),
]
x = Inches(0.62)
for t, d, c in subs:
    w = Inches(2.98)
    add_rect(s, x, Inches(3.2), w, Inches(1.25), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(3.2), w, Inches(0.1), c)
    add_text(s, x + Inches(0.15), Inches(3.42), w - Inches(0.3), Inches(0.42),
             t, size=14, color=c, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.15), Inches(3.86), w - Inches(0.3), Inches(0.5),
             d, size=10.5, color=DARK, align=PP_ALIGN.CENTER)
    x += w + Inches(0.12)
add_text(s, Inches(0.62), Inches(4.72), Inches(12.1), Inches(0.4),
         "仪式感与符号系统(把「起点」做成可带走的产品):", size=13, color=PRIMARY_D, bold=True)
rite = [
    ("零公里发车仪式", "每日固定时段发车礼:鸣笛、盖章、合影,拍摄位标准化"),
    ("环游中国第一枚章", "「起点护照」集章体系,串联全市景点与沿线城市"),
    ("起点纪念物系", "车贴/邮戳/完驾证书/银杏书签,低成本高传播的社交货币"),
    ("事件日历", "自驾集结赛 · 鸭绿江马拉松 · 银杏文化周 · 开海节四季不断"),
]
x = Inches(0.62)
for t, d in rite:
    w = Inches(2.98)
    add_rect(s, x, Inches(5.18), w, Inches(1.35), None, line=LINEGRAY,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.15), Inches(5.32), w - Inches(0.3), Inches(0.4),
             t, size=12, color=ACCENT, bold=True)
    add_text(s, x + Inches(0.15), Inches(5.74), w - Inches(0.3), Inches(0.75),
             d, size=10, color=DARK, line_spacing=1.15)
    x += w + Inches(0.12)

# S: 沿江
s = new_slide(); set_bg(s)
header(s, "PART 04 · 策略主张", "沿江品质提升:虎山—断桥「一带四节点」",
       "约20公里黄金岸线,承担游客第一印象的「中央体验区」")
nodes = [
    ("① 虎山长城", "起点仪式区", "长城东端起点+零公里联动;发车仪式主场;夜游虎山长城", TEAL),
    ("② 国门湾", "营地集散区", "房车营地扩容连锁;自驾服务中心(补给/维保/攻略);观江平台", PRIMARY),
    ("③ 老城滨江", "生活休闲区", "滨江慢行道贯通;安东老街扩容;江畔市集与银杏大道", ACCENT),
    ("④ 断桥核心", "沉浸体验区", "桥体光影夜游;1950沉浸演艺;纪念馆—断桥—老街一体动线", RED),
]
x = Inches(0.62)
for i, (t, tag, d, c) in enumerate(nodes):
    w = Inches(2.85)
    add_rect(s, x, Inches(1.8), w, Inches(2.5), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.8), w, Inches(0.12), c)
    add_text(s, x + Inches(0.16), Inches(2.05), w - Inches(0.32), Inches(0.42),
             t, size=14, color=c, bold=True)
    add_text(s, x + Inches(0.16), Inches(2.5), w - Inches(0.32), Inches(0.35),
             tag, size=11, color=PRIMARY_D, bold=True)
    add_text(s, x + Inches(0.16), Inches(2.9), w - Inches(0.32), Inches(1.3),
             d, size=10, color=DARK, line_spacing=1.2)
    if i < 3:
        ar = slide_arrow = add_rect(s, x + w + Inches(0.02), Inches(2.85), Inches(0.24), Inches(0.4),
                                    ACCENT, shape=MSO_SHAPE.RIGHT_ARROW)
    x += w + Inches(0.28)
bullets(s, Inches(0.62), Inches(4.62), Inches(12.1), Inches(1.9), [
    ("慢行系统:", "滨江骑行道/步道全线贯通,驿站(补给+卫生间+充电)按2–3公里间距布点,标识统一到「起点」视觉体系。"),
    ("游船升级:", "分层产品——跨境视野观光航班/日落航班/餐饮演艺主题航班;探索冬季航线,配合「夜航+夜市+夜演」三夜联动。"),
    ("城市更新结合:", "滨江老旧街区与轻纺工业遗存「微更新+运营前置」,改造为民宿、市集、展演空间,避免大拆大建。"),
], size=11.5, gap=6)

# S: 联动
s = new_slide(); set_bg(s)
header(s, "PART 04 · 策略主张", "全域联动:「一江两线三圈」",
       "让游客的下一步永远有去处——把散点资源编织成连续动线")
add_rect(s, Inches(0.62), Inches(1.75), Inches(5.6), Inches(0.9), PRIMARY_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(0.85), Inches(1.75), Inches(5.2), Inches(0.9),
         [([("一江:", {"color": ACCENT, "bold": True}), ("鸭绿江城市段主轴(中央体验区)", {"color": WHITE})], {})],
         size=13, anchor=MSO_ANCHOR.MIDDLE)
lines = [
    ("两线 · G331边境线", "虎山 → 河口(桃花/毛岸英学校) → 绿江村(塞外小江南),边境自驾景观廊道", TEAL),
    ("两线 · G228滨海线", "东港(海鲜/湿地观鸟) → 大鹿岛 → 獐岛(海疆最北端起点),滨海度假走廊", PRIMARY),
]
y = Inches(2.8)
for t, d, c in lines:
    add_rect(s, Inches(0.62), y, Inches(5.6), Inches(0.95), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, Inches(0.62), y, Inches(0.08), Inches(0.95), c)
    add_text(s, Inches(0.85), y + Inches(0.08), Inches(5.2), Inches(0.38), t, size=12, color=c, bold=True)
    add_text(s, Inches(0.85), y + Inches(0.46), Inches(5.2), Inches(0.45), d, size=10, color=DARK)
    y += Inches(1.05)
circles = [
    ("城市核心圈", "断桥/纪念馆/锦江山/安东老街/银杏街", RED),
    ("宽甸山水圈", "天桥沟/青山沟/黄椅山/温泉带", TEAL),
    ("东港海岛圈", "大鹿岛/獐岛/孤山镇/海鲜市集", PRIMARY),
]
x = Inches(6.5)
for t, d, c in circles:
    add_rect(s, x, Inches(1.75), Inches(2.0), Inches(2.0), None, line=c, shape=MSO_SHAPE.OVAL)
    add_text(s, x + Inches(0.15), Inches(2.25), Inches(1.7), Inches(0.4),
             t, size=12, color=c, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.68), Inches(1.6), Inches(0.9),
             d, size=8.5, color=DARK, align=PP_ALIGN.CENTER, line_spacing=1.1)
    x += Inches(2.15)
bullets(s, Inches(6.5), Inches(4.05), Inches(6.3), Inches(1.85), [
    ("交通组织:", "旅游环线巴士+景区接驳专线;「高铁+落地租车/房车」产品包;营地网络按线布点。"),
    ("区域结盟:", "发起G331沿线城市文旅联盟(丹东—通化—白山—延边),互认通票、共推「一程多站」。"),
    ("铁路联动:", "深度接入东北「旅游计次票」体系(已覆盖丹东站),承接哈尔滨/长白山溢出客流。"),
], size=11, gap=6)
note_bar(s, Inches(0.62), Inches(6.15), Inches(12.1),
         "联动目标:把「半日打卡城市」变成「2晚3天目的地」——每增加一晚停留,人均消费预计提升60%以上。")

# S: 票务
s = new_slide(); set_bg(s)
header(s, "PART 04 · 策略主张", "票务组合:低门槛进入,高体验变现",
       "借鉴万岁山「低价高频」逻辑,重构「票务让利、二消挣钱」的收入模型")
tks = [
    ("丹东City Pass城市通票", "3日/5日两档,通行主要A级景区+游船/温泉折扣+环线巴士免费;定价锚定「单景区两次门票之和」以下", PRIMARY),
    ("「起点护照」联票", "集章免费、权益付费:集满兑换纪念品/餐饮券,绑定复游;与G331沿线城市互认", ACCENT),
    ("一票三日无限次试点", "选取1–2个核心景区(如虎山长城)试点,验证「长停留带动周边二消」模型", TEAL),
    ("旅居者年卡", "面向银发旅居与本地市民的低价年卡,平峰期填舱,培养「城市主人」口碑", RED),
]
x, y0 = Inches(0.62), Inches(1.78)
for i, (t, d, c) in enumerate(tks):
    cx = x + (i % 2) * Inches(6.2)
    cy = y0 + (i // 2) * Inches(1.6)
    add_rect(s, cx, cy, Inches(5.95), Inches(1.45), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, cx, cy, Inches(0.08), Inches(1.45), c)
    add_text(s, cx + Inches(0.25), cy + Inches(0.12), Inches(5.5), Inches(0.4),
             t, size=13, color=c, bold=True)
    add_text(s, cx + Inches(0.25), cy + Inches(0.55), Inches(5.5), Inches(0.8),
             d, size=10.5, color=DARK, line_spacing=1.2)
add_text(s, Inches(0.62), Inches(5.15), Inches(12.1), Inches(0.4),
         "配套原则:", size=13, color=PRIMARY_D, bold=True)
bullets(s, Inches(0.62), Inches(5.55), Inches(12.1), Inches(1.3), [
    ("极限淡季定价:", "冬春季通票价格下探至旺季1/3,以「极致性价比旅居」对冲季节性(参考东北铁路计次票逻辑)。"),
    ("收入再平衡:", "门票收入占比目标从主导降至30%以下,住宿/餐饮/演艺/文创成为主力,倒逼二消场景建设。"),
], size=11.5, gap=6)

# S: 产品组合
s = new_slide(); set_bg(s)
header(s, "PART 04 · 策略主张", "产品组合:自驾优先,四季不断档",
       "六大产品线对应六类客群,按投入强度分期推进")
rows = [
    ["产品线", "核心内容", "目标客群", "优先级"],
    ["自驾/公路旅行", "G331首发站产品包(发车仪式+车检补给+首晚住宿+攻略);房车营地连锁;摩旅/骑行友好配套", "全国自驾圈层", "★★★ 近期"],
    ["红色研学", "纪念馆+断桥+河口标准化课程;1950沉浸式演艺;国防教育营", "学校/机关/亲子", "★★★ 近期"],
    ["康养旅居", "「温泉+海岛+江景」月付套餐;候鸟公寓;医养配套", "银发/候鸟客群", "★★★ 近期"],
    ["边境体验", "跨境视野游船;国门与界碑打卡;朝鲜族民俗村与服饰写真;边贸市集", "年轻圈层/亲子", "★★ 中期"],
    ["美食物产", "黄蚬子/开海节事件化;草莓季(采摘+甜品节);伴手礼体系", "全客群", "★★ 中期"],
    ["沉浸内容", "安东老街NPC常态化;鸭绿江夜游演艺;冬季冰凌+温泉产品", "年轻圈层", "★ 中远期"],
]
add_table(s, Inches(0.55), Inches(1.7), Inches(12.25), Inches(4.15), rows,
          col_widths=[1.1, 3.3, 1.2, 0.95], font_size=10.5, header_size=11.5)
note_bar(s, Inches(0.55), Inches(6.1), Inches(12.25),
         "四季日历:春(河口桃花/草莓尾季) → 夏(避暑+海岛+自驾旺季) → 秋(银杏+枫叶双景) → 冬(温泉+冰凌+极限性价比旅居),用产品填平淡季而非用折扣硬扛。")

# ============================================================
# PART 05  客群导入与更高维度
# ============================================================
section_slide("05", "客群导入与更高维度价值",
              ["客群分层:基本盘 → 增量主力 → 战略客群的导入次序",
               "银发经济:从旅游子项升格为城市战略",
               "旅游作为杠杆:人口回流 · 城市更新 · 产业联动"])

# S: 客群导入
s = new_slide(); set_bg(s)
header(s, "PART 05 · 客群与更高维度", "客群导入:分层次序与转化机制",
       "先守基本盘,再攻增量,战略客群决定城市长期曲线")
tiers = [
    ("基本盘", "辽吉黑都市圈短途周末客群", "沈阳/大连2小时圈;周末游+美食游;复游率是核心指标", PRIMARY, Inches(11.4)),
    ("增量主力", "京津冀高铁客群 + 全国自驾/摩旅圈层", "高铁5小时圈专列/专题;自驾俱乐部B端合作;首发站产品承接", ACCENT, Inches(9.6)),
    ("战略客群", "银发旅居候鸟 + 亲子研学", "康养机构/老年大学直连;研学课程共建;按月旅居产品", TEAL, Inches(7.8)),
    ("远期客群", "入境边境游(政策窗口)", "中韩班轮复航跟踪;边境旅游试验区申创;跨境线路储备", RED, Inches(6.0)),
]
y = Inches(1.78)
for t, who, how, c, w in tiers:
    add_rect(s, Inches(0.62), y, w, Inches(1.06), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, Inches(0.62), y, Inches(0.08), Inches(1.06), c)
    add_text(s, Inches(0.88), y + Inches(0.1), Inches(1.3), Inches(0.8),
             t, size=13, color=c, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.25), y + Inches(0.1), w - Inches(1.9), Inches(0.4), who, size=12, color=PRIMARY_D, bold=True)
    add_text(s, Inches(2.25), y + Inches(0.52), w - Inches(1.9), Inches(0.45), how, size=10.5, color=DARK)
    y += Inches(1.16)
add_rect(s, Inches(9.0), Inches(3.0), Inches(3.7), Inches(3.3), PRIMARY_D, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(s, Inches(9.25), Inches(3.2), Inches(3.2), Inches(0.4),
         "流量 → 留量 → 复游", size=14, color=ACCENT, bold=True)
add_text(s, Inches(9.25), Inches(3.66), Inches(3.25), Inches(2.5),
         [("① 延长停留:夜产品+通票+营地", {"space_after": 7}),
          ("② 提高客单:二消场景+体验溢价", {"space_after": 7}),
          ("③ 制造复游:集章权益+季节限定+旅居身份", {"space_after": 7}),
          ("④ 放大口碑:服务兜底+UGC激励", {})],
         size=11, color=WHITE, line_spacing=1.25)

# S: 银发经济
s = new_slide(); set_bg(s)
header(s, "PART 05 · 客群与更高维度", "银发经济:丹东的「第二人口」战略",
       "用旅居人口对冲户籍人口流出——淡季、房产、就业三个问题一并求解")
cols = [
    ("为什么是丹东", [
        "夏季凉爽宜居,是南方「候鸟」避暑的天然目的地",
        "五龙背/东汤温泉+江海景观,康养资源组合东北少有",
        "房价与生活成本低,存量住宅充裕,旅居性价比全国前列",
        "医疗基础尚可,毗邻沈阳大连优质医疗圈",
    ], TEAL),
    ("产品与供给", [
        "「温泉+海岛+江景」月付旅居套餐,按30/60/90天设计",
        "存量住宅改造候鸟公寓,统一托管运营",
        "旅居者市民待遇包:公交/图书馆/社区医疗/老年大学",
        "医养结合:体检、慢病管理、紧急救护网络",
    ], PRIMARY),
    ("城市收益", [
        "平峰季稳定现金流,「淡季不淡」结构性改善",
        "盘活闲置房产,稳定物业与社区服务就业",
        "旅居者成为口碑节点,带动亲友探访式旅游",
        "部分旅居转定居,直接贡献人口回流",
    ], ACCENT),
]
x = Inches(0.62)
for t, items, c in cols:
    w = Inches(3.95)
    add_rect(s, x, Inches(1.75), w, Inches(4.35), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(s, x, Inches(1.75), w, Inches(0.12), c)
    add_text(s, x + Inches(0.2), Inches(2.0), w - Inches(0.4), Inches(0.45),
             t, size=14.5, color=c, bold=True)
    add_text(s, x + Inches(0.2), Inches(2.5), w - Inches(0.4), Inches(3.4),
             [([("· ", {"color": c, "bold": True}), (it, {})], {"space_after": 8, "line_spacing": 1.18}) for it in items],
             size=10.5, color=DARK)
    x += w + Inches(0.12)
note_bar(s, Inches(0.62), Inches(6.3), Inches(12.1),
         "银发旅居的本质是「把人留下来生活」——它对城市GDP、就业与人口结构的贡献,可能超过观光旅游本身。")

# S: 人口回流与城市更新
s = new_slide(); set_bg(s)
header(s, "PART 05 · 客群与更高维度", "人口回流与城市更新:旅游作为城市杠杆",
       "文旅不只是产业,更是重新配置城市资产与人口的抓手")
lcol = [
    ("文旅创业生态", "民宿/营地/餐饮/文创/内容创作者的小微创业扶持基金与轻审批;「青年返乡合伙人」计划,优先给流量位与租金减免"),
    ("就业带动", "自驾服务、康养照护、演艺文创等新岗位定向培训;文旅从业者技能认证与体面薪酬示范"),
    ("新丹东人政策", "旅居转定居的落户/购房/医保衔接便利;引进运营人才「一事一议」"),
]
rcol = [
    ("滨江带更新", "老城滨江生活岸线贯通,工业遗存(轻纺记忆)改造为展演/市集/民宿空间,「微更新+运营前置」"),
    ("老街区文旅化", "安东老城肌理保护性利用,避免仿古一条街同质化;鼓励原住民业态参与,保留生活气"),
    ("闲置资产盘活", "闲置厂房/楼宇转文旅新空间白名单制度;存量住宅→候鸟公寓的合规改造通道"),
]
add_text(s, Inches(0.62), Inches(1.7), Inches(5.9), Inches(0.45),
         "人口回流:让旅游岗位接住返乡青年", size=15, color=PRIMARY_D, bold=True)
y = Inches(2.2)
for t, d in lcol:
    add_rect(s, Inches(0.62), y, Inches(5.9), Inches(1.28), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(0.85), y + Inches(0.1), Inches(5.45), Inches(0.4), t, size=12.5, color=TEAL, bold=True)
    add_text(s, Inches(0.85), y + Inches(0.5), Inches(5.45), Inches(0.72), d, size=10.5, color=DARK, line_spacing=1.15)
    y += Inches(1.42)
add_text(s, Inches(6.85), Inches(1.7), Inches(5.9), Inches(0.45),
         "城市更新:让旅游流量反哺城市资产", size=15, color=PRIMARY_D, bold=True)
y = Inches(2.2)
for t, d in rcol:
    add_rect(s, Inches(6.85), y, Inches(5.9), Inches(1.28), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, Inches(7.08), y + Inches(0.1), Inches(5.45), Inches(0.4), t, size=12.5, color=RED, bold=True)
    add_text(s, Inches(7.08), y + Inches(0.5), Inches(5.45), Inches(0.72), d, size=10.5, color=DARK, line_spacing=1.15)
    y += Inches(1.42)
note_bar(s, Inches(0.62), Inches(6.5), Inches(12.1),
         "产业联动:农文旅(草莓/板栗/海产深加工伴手礼)+边贸购物+枢纽经济(高铁+双国道),让旅游成为多产业的流量入口。", h=Inches(0.52))

# ============================================================
# PART 06  实施路径
# ============================================================
section_slide("06", "实施路径与保障",
              ["三阶段推进:速赢 → 系统 → 深化",
               "机制:平台公司+专业运营团队+民营业态生态",
               "风险边界:政策敏感性 · 叙事庄重底线 · 季节现金流"])

# S: 三阶段
s = new_slide(); set_bg(s)
header(s, "PART 06 · 实施路径", "三阶段路线图:先声量,后体系,再旗舰", None)
phases = [
    ("第一阶段 · 速赢起势", "0–12个月", [
        "零公里仪式系统+起点护照上线",
        "City Pass通票与淡季极限定价",
        "沿江驿站/标识/慢行道补短板",
        "事件日历落地(自驾集结/马拉松/银杏周/开海节)",
        "「不让游客受委屈」服务承诺与先行赔付",
    ], ACCENT),
    ("第二阶段 · 系统建设", "1–3年", [
        "虎山—断桥「一带四节点」整体运营",
        "房车营地网络与自驾服务中心连锁",
        "银发旅居产品与候鸟公寓规模化",
        "安东老街NPC演艺常态化+夜航夜市夜演",
        "G331沿线城市联盟与通票互认",
    ], PRIMARY),
    ("第三阶段 · 旗舰深化", "3–5年", [
        "鸭绿江夜游演艺/1950沉浸剧场旗舰项目",
        "创建国家边境旅游试验区,入境产品上线",
        "滨江片区级城市更新",
        "冲击5A级景区与国家级旅游度假区",
        "「旅居转定居」人口政策闭环",
    ], RED),
]
x = Inches(0.62)
for t, tm, items, c in phases:
    w = Inches(3.95)
    add_rect(s, x, Inches(1.7), w, Inches(0.85), c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.2), Inches(1.78), w - Inches(0.4), Inches(0.4), t, size=13.5, color=WHITE, bold=True)
    add_text(s, x + Inches(0.2), Inches(2.16), w - Inches(0.4), Inches(0.32), tm, size=10.5,
             color=RGBColor(0xF2, 0xF5, 0xFA))
    add_rect(s, x, Inches(2.68), w, Inches(3.6), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(s, x + Inches(0.2), Inches(2.85), w - Inches(0.4), Inches(3.3),
             [([("· ", {"color": c, "bold": True}), (it, {})], {"space_after": 8, "line_spacing": 1.15}) for it in items],
             size=10.5, color=DARK)
    x += w + Inches(0.12)

# S: 机制与风险
s = new_slide(); set_bg(s)
header(s, "PART 06 · 实施路径", "机制保障与风险边界", None)
add_text(s, Inches(0.62), Inches(1.68), Inches(5.9), Inches(0.42),
         "运营机制", size=15, color=PRIMARY_D, bold=True)
bullets(s, Inches(0.62), Inches(2.15), Inches(5.9), Inches(4.0), [
    ("政企分工:", "管委会管资源与规则,平台公司管资产,专业运营团队管产品与内容——学河南引入市场化运营主体。"),
    ("民营生态:", "餐饮/民宿/营地/文创放开给民营与个体,政府做基础设施与流量分发,不与民争利。"),
    ("数据公示:", "客流/收入/投诉数据定期公开,建立经营信任(参考万岁山做法)。"),
    ("全民动员:", "「人人都是丹东城市形象大使」;出租车/餐饮/住宿服务规范与暗访机制。"),
], size=12, gap=9)
add_text(s, Inches(6.85), Inches(1.68), Inches(5.9), Inches(0.42),
         "风险与边界", size=15, color=RED, bold=True)
bullets(s, Inches(6.85), Inches(2.15), Inches(5.9), Inches(4.0), [
    ("涉朝敏感性:", "跨境元素严格限定在合规观光范畴;入境/跨境产品跟随国家政策窗口,不抢跑。"),
    ("红色叙事底线:", "抗美援朝主题坚持庄重表达,严禁过度娱乐化;沉浸式≠游乐化。"),
    ("季节现金流:", "重资产项目谨慎上马,优先轻资产内容与存量改造;用旅居与冬季产品平滑现金流。"),
    ("同质化风险:", "拒绝「仿古街+灯光秀」通用模板,所有产品必须回答「为什么只能在丹东」。"),
], size=12, gap=9)

# S: 结束页
s = new_slide(); set_bg(s, PRIMARY_D)
PAGE_NO[0] += 1
add_rect(s, 0, 0, SW, Inches(0.18), ACCENT)
add_rect(s, 0, Inches(7.32), SW, Inches(0.18), ACCENT)
add_text(s, Inches(0.95), Inches(2.3), Inches(11.4), Inches(1.0),
         "从流量到留量,从起点到目的地", size=34, color=WHITE, bold=True)
add_text(s, Inches(0.95), Inches(3.45), Inches(11.4), Inches(1.6),
         [("丹东不需要成为第二个哈尔滨——", {"space_after": 6}),
          ("它只需要成为中国边境旅行无可替代的第一站,和让人愿意住下来的暖城。", {})],
         size=16, color=RGBColor(0xC7, 0xD6, 0xE8), line_spacing=1.4)
add_text(s, Inches(0.95), Inches(5.3), Inches(11.4), Inches(0.9),
         [("下一步:确认本报告框架与观点后,可深化为——", {"space_after": 5}),
          ("① 分项行动清单与责任分工建议    ② 重点项目概念方案    ③ 投入产出与客流测算", {})],
         size=12.5, color=ACCENT, line_spacing=1.35)
add_text(s, Inches(0.95), Inches(6.6), Inches(11), Inches(0.4),
         "鸭绿江畔 · 丹东真好", size=13, color=RGBColor(0x8F, 0xA6, 0xC0))

OUT = "丹东城市旅游价值提升策略报告.pptx"
prs.save(OUT)
print(f"saved: {OUT}, slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
